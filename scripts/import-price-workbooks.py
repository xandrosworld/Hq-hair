import openpyxl,json,re,hashlib,sys
from pathlib import Path
from decimal import Decimal,ROUND_HALF_UP
root=Path(sys.argv[1]) if len(sys.argv)>1 else Path(r'C:\Users\DELL\Downloads\ai');prices=[];coverage=[]
for file,tier in [('Bảng giá Basic so9 (1).xlsx','Basic'),('Gía Baby so9.xlsx','Baby'),('Premium HQ so 9.xlsx','Premium')]:
 w=openpyxl.load_workbook(root/file,data_only=True)
 for s in w:
  count=len(prices);used=set()
  def block(start,end,ic,cc,pcs,variant='',segment=''):
   for r in range(start,end+1):
    inch=s.cell(r,ic).value;cm=s.cell(r,cc).value
    if not isinstance(inch,(int,float)) or not isinstance(cm,(int,float)):continue
    used.update([s.cell(r,ic).coordinate,s.cell(r,cc).coordinate])
    for col,tone in pcs:
     c=s.cell(r,col);v=c.value
     if not isinstance(v,(int,float)):continue
     assert c.coordinate not in used,(s.title,c.coordinate)
     used.add(c.coordinate);ident=f'{tier}:{s.title}:{c.coordinate}'
     prices.append(dict(id=hashlib.sha256(ident.encode()).hexdigest()[:20],tier=tier,product=s.title,variant=variant,segment=segment,inches=inch,lengthCm=cm,tone=tone,price=float(Decimal(str(v)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)),priceBasis='unit' if any(x in s.title.lower() for x in ['closure','frontal','ponytail','topper']) else '100g',source={'file':file,'sheet':s.title,'cell':c.coordinate}))
  trip=lambda c:[(c,'black'),(c+1,'brown'),(c+2,'blonde')]
  if s.title=='Closure':
   for a,b,i,c,v in [(4,20,1,3,'2x4'),(4,20,9,11,'2x6'),(5,18,18,20,'4x4'),(21,33,18,20,'5x5'),(25,35,1,3,'6x6'),(26,36,9,11,'7x7')]:block(a,b,i,i+1,trip(c),v)
  elif s.title in ['Frontal','HD Lace']:
   for row in s:
    for c in row:
     if isinstance(c.value,str) and c.value.lower().strip()=='inches':
      r=c.row;col=c.column;start=col+2
      titles=[s.cell(r-1,j).value for j in range(col,col+5) if isinstance(s.cell(r-1,j).value,str)]
      variant=titles[0] if titles else ''
      end=r+1
      while isinstance(s.cell(end,col).value,(int,float)):end+=1
      block(r+1,end-1,col,col+1,trip(start),variant)
  elif s.title=='Topper full lưới':
   block(4,12,3,4,trip(5),'5x5','2D');block(15,23,3,4,trip(5),'5x5','3D')
  elif s.title=='Topper Mix Weft':
   for hr,seg in [(6,'2D'),(30,'3D')]:
    for ic,pc,tone in [(1,3,'brown'),(12,14,'black'),(22,24,'blonde')]:
     for c in range(pc,pc+5):block(hr+1,hr+12,ic,ic+1,[(c,tone)],str(s.cell(hr,c).value),seg)
  elif s.title=='Raw bao tẩy':block(1,s.max_row,1,2,[(3,'black')])
  else:
   double=any('double drawn' in str(c.value).lower() for row in s for c in row)
   if double:block(1,s.max_row,1,2,trip(3),segment='Double Drawn');block(1,s.max_row,1,2,trip(6),segment='Super Double Drawn')
   else:block(1,s.max_row,1,2,trip(3))
  # Every numeric non-length cell must be accounted for, including staggered table blocks.
  missing=[]
  for row in s:
   for c in row:
    if isinstance(c.value,(int,float)) and c.coordinate not in used:
     missing.append(c.coordinate)
  coverage.append(dict(tier=tier,sheet=s.title,prices=len(prices)-count,unmapped=missing))
colors=[];data=json.loads(Path('data/imports/colors-public.json').read_text())['data'];cols={c['name']:c for c in data['tableSchemas'][0]['columns']};choices=cols['Category']['typeOptions']['choices']
for r in data['tableDatas'][0]['rows']:
 v=r['cellValuesByColumnId'];category=choices[v[cols['Category']['id']]]['name'];code=v[cols['Color Code']['id']].strip();images=[]
 for field in ['Color Chart','Real Images']:
  for a in v.get(cols[field]['id'],[]):
   if a.get('type','').startswith('image/'):images.append(dict(id=a['id'],kind=field))
 colors.append(dict(id=r['id'],code=code,category=category,tone={'Black':'black','Brown':'brown','Blonde':'blonde','Other':'blonde'}.get(category),components=[],images=images,note=v.get(cols['Note']['id'],'')))
base={c['code'].lstrip('#').upper():c for c in colors if c['tone']}
rank={'black':0,'brown':1,'blonde':2}
for c in colors:
 c['sourceUrl']='https://airtable.com/appwCEKi3pHQ4CFEB/shrOpuqhphiyQyyQ8/tblOuyQezeDc71KAw/viwoZzKUus5StmDHt'
 record=next(r for r in data['tableDatas'][0]['rows'] if r['id']==c['id'])
 c['videoCount']=sum(a.get('type','').startswith('video/') for field in ['Color Chart','Real Images'] for a in record['cellValuesByColumnId'].get(cols[field]['id'],[]))
 if not c['tone']:
  parts=re.split(r'[-/]',re.sub(r'^(Ombre|Piano|Balayage)\s*','',c['code'].lstrip('#'),flags=re.I));parts=[p.strip().upper() for p in parts]
  c['components']=[base[p]['id'] for p in parts if p in base];unknown=[p for p in parts if p not in base]
  # A known blonde component already determines the highest tone even if another is unlisted.
  tones=[base[p]['tone'] for p in parts if p in base]
  c['tone']='blonde' if 'blonde' in tones else None if unknown else max(tones,key=rank.get)
  c['unlistedComponents']=unknown
assert not any(c['unmapped'] for c in coverage),'Unmapped numeric cells: inspect coverage before import'
assert len(colors)==94,'Unexpected source color count'
out={'version':1,'prices':prices,'colors':colors,'sourceDate':'2026-09-30'}
Path('resources/pricing-seed.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
Path('data/imports/price-coverage.json').write_text(json.dumps(coverage,ensure_ascii=False,indent=2),encoding='utf-8')
print('PRICES',len(prices),'COLORS',len(colors));print('UNMAPPED',[c for c in coverage if c['unmapped']]);print('UNRESOLVED',[(c['code'],c.get('unlistedComponents')) for c in colors if not c['tone']]);print('COUNTS',[(t,sum(p['tier']==t for p in prices)) for t in ['Basic','Baby','Premium']])
